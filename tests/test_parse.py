"""The converters of drakkar-parse, on small samples of each file type (no network)."""

import json
import subprocess
from datetime import datetime
from pathlib import Path

import docx
import openpyxl
import pymupdf
import pytest
import xlwt

from drakkar.parse import ParseError, convert, markdown_table, parse, sheets

# ppat.1002331.s00*.doc: supplementary tables of PMID 22046132 (CC BY). Ours, made once with
# LibreOffice from a .docx (LibreOffice is no dependency): beyond-bmp.doc, and page-break.doc,
# whose first paragraph mark was then turned into a page break, as Word writes one.
SAMPLES = Path(__file__).parent / "samples"

JATS = """<?xml version="1.0"?>
<article xmlns:xlink="http://www.w3.org/1999/xlink">
  <front><article-meta>
    <title-group><article-title>NS1 binds <italic>TRIM25</italic></article-title></title-group>
    <abstract><p>We show that NS1 binds TRIM25.</p></abstract>
  </article-meta></front>
  <body>
    <sec><title>Results</title>
      <p>NS1 co-precipitates TRIM25 (<xref ref-type="fig" rid="f1">Fig. 1</xref>).
        <table-wrap><label>Table 1</label><caption><p>Interactions.</p></caption>
          <table>
            <thead><tr><th>Bait</th><th colspan="2">Prey</th></tr></thead>
            <tbody>
              <tr><td rowspan="2">NS1</td><td>TRIM25</td><td>Q14258</td></tr>
              <tr><td>RIG-I</td><td>O95786</td></tr>
            </tbody>
          </table>
        </table-wrap>
      </p>
      <fig id="f1"><label>Figure 1</label>
        <caption><title>Co-IP.</title><p>Anti-FLAG.</p></caption></fig>
      <sec><title>Mapping</title><p>Residues 1-73 of NS1.</p></sec>
    </sec>
  </body>
  <back>
    <ack><p>We thank everyone.</p></ack>
    <ref-list><ref><mixed-citation>A cited paper.</mixed-citation></ref></ref-list>
    <sec><title>Supporting Information</title>
      <supplementary-material xlink:href="s1.xlsx"><label>Table S1</label>
        <caption><p>Y2H screen.</p></caption></supplementary-material>
    </sec>
  </back>
</article>
"""


def test_jats(tmp_path: Path):
    path = tmp_path / "a.xml"
    path.write_text(JATS)
    assert parse(path) == (
        "# NS1 binds TRIM25\n\n"
        "## Abstract\n\nWe show that NS1 binds TRIM25.\n\n"
        "## Results\n\n"
        "NS1 co-precipitates TRIM25 (Fig. 1).\n\n"
        "**Table 1.** Interactions.\n\n"
        "| Bait | Prey | Prey |\n|---|---|---|\n"
        "| NS1 | TRIM25 | Q14258 |\n| NS1 | RIG-I | O95786 |\n\n"
        "**Figure 1.** Co-IP. Anti-FLAG.\n\n"
        "### Mapping\n\nResidues 1-73 of NS1.\n\n"
        "## Supporting Information\n\n"
        "**Table S1.** Y2H screen. (file s1.xlsx)\n"
    )


def test_jats_formula_alternatives_and_bad_spans(tmp_path: Path):
    path = tmp_path / "a.xml"
    path.write_text(
        "<article><body><p>K<sub>d</sub> = <inline-formula><alternatives>"
        "<tex-math>10^{-6}</tex-math><mml:math xmlns:mml='http://www.w3.org/1998/Math/MathML'>"
        "<mml:mn>10</mml:mn></mml:math></alternatives></inline-formula> M</p>"
        "<table-wrap><table><tr><td colspan='x'>a</td><td>b</td></tr></table></table-wrap>"
        "</body></article>"
    )
    assert parse(path) == "Kd = 10^{-6} M\n\n**Table.**\n\n| a | b |\n|---|---|\n"


def test_jats_block_alternatives(tmp_path: Path):
    path = tmp_path / "a.xml"
    path.write_text(
        "<article><body><p>The ratio:</p><disp-formula><alternatives><tex-math>a/b</tex-math>"
        "<mml:math xmlns:mml='http://www.w3.org/1998/Math/MathML'><mml:mi>a</mml:mi></mml:math>"
        "</alternatives></disp-formula><alternatives><p>First notation.</p><p>Second.</p>"
        "</alternatives></body></article>"
    )
    assert parse(path) == "The ratio:\n\na/b\n\nFirst notation.\n"


def test_jats_namespaced(tmp_path: Path):
    path = tmp_path / "a.xml"
    path.write_text(
        '<article xmlns="http://jats.nlm.nih.gov"><front><article-meta><title-group>'
        "<article-title>NS1 binds TRIM25</article-title></title-group></article-meta></front>"
        "<body><sec><title>Results</title><p>Co-IP.</p></sec></body></article>"
    )
    assert parse(path) == "# NS1 binds TRIM25\n\n## Results\n\nCo-IP.\n"


def test_jats_text_inside_paragraphs(tmp_path: Path):
    """In PMC XML, the text of list items, definitions, table cells and table footnotes sits in
    paragraphs of its own."""
    path = tmp_path / "a.xml"
    path.write_text(
        "<article><body><sec><title>Methods</title>"
        "<list><list-item><p>Y2H screen</p><list><list-item><p>ORFeome</p></list-item></list>"
        "</list-item><list-item><p>LUMIER</p></list-item></list>"
        "<glossary><def-list><def-item><term>IP</term><def><p>immunoprecipitation</p></def>"
        "</def-item></def-list></glossary>"
        "<table-wrap><table><tr><td><p>NS1</p></td><td><p>TRIM25<sup>a</sup></p></td></tr>"
        "</table><table-wrap-foot><fn><p><sup>a</sup> Confirmed by co-IP.</p></fn>"
        "</table-wrap-foot></table-wrap></sec></body></article>"
    )
    assert parse(path) == (
        "## Methods\n\n"
        "- Y2H screen\n- ORFeome\n- LUMIER\n\n"
        "- IP immunoprecipitation\n\n"
        "**Table.**\n\n| NS1 | TRIM25a |\n|---|---|\n\n"
        "a Confirmed by co-IP.\n"
    )


def test_markdown_table_escapes_pipes_and_pads_rows():
    assert markdown_table([["a", "b|c"], ["d"]]) == "| a | b\\|c |\n|---|---|\n| d |  |"


def test_html(tmp_path: Path):
    paragraphs = "".join(
        f"<p>Paragraph {i}: NS1 co-precipitates TRIM25 in HEK293 cells, as the blots show.</p>"
        for i in range(30)
    )
    path = tmp_path / "a.html"
    path.write_text(
        "<html><head><title>T</title><script>var x = 1;</script></head><body>"
        "<nav><a href='/'>Home</a></nav>"
        f"<article><h1>NS1 binds TRIM25</h1><h2>Results</h2>{paragraphs}"
        "<table><tr><th>Bait</th><th>Prey</th></tr><tr><td>NS1</td><td>TRIM25</td></tr></table>"
        "</article></body></html>"
    )
    text = parse(path)
    assert "Paragraph 29: NS1 co-precipitates TRIM25" in text
    assert "| NS1 | TRIM25 |" in text
    assert "var x" not in text


PMC_PAGE = """<html><body><hgroup><h1>NS1 binds <i>TRIM25</i></h1></hgroup>
<section class="body main-article-body">
  <section class="abstract" id="abstract1"><h2 class="pmc_sec_title">Abstract</h2>
    <p>We show that NS1 binds TRIM25.</p></section>
  <section id="kwd-group1" class="kwd-group"><p>influenza</p></section>
  <section id="S1"><h2 class="pmc_sec_title">Results</h2>
    <p>NS1 co-precipitates TRIM25 (<a href="#F1">Fig. 1</a>).</p>
    <figure class="fig xbox font-sm" id="F1"><h3 class="obj_head">Figure 1.</h3>
      <p class="img-box"><img src="f1.jpg"></p><figcaption><p>Co-IP. Anti-FLAG.</p></figcaption>
    </figure>
    <section class="tw xbox font-sm" id="T1"><h4 class="obj_head">Table 1.</h4>
      <div class="caption p"><p>Interactions.</p></div>
      <div class="tbl-box p"><table><thead><tr><th>Bait</th><th colspan="2">Prey</th></tr></thead>
        <tbody><tr><td>NS1</td><td>TRIM25</td><td>Q14258</td></tr></tbody></table></div>
      <div class="tw-foot p"><div class="fn"><p>Co-IP in HEK293T cells.</p></div></div>
    </section>
    <section id="S1.1"><h3 class="pmc_sec_title">Mapping</h3>
      <ul><li>Residues 1-73 of NS1.</li></ul></section>
  </section>
  <section class="ack" id="ack1"><h2 class="pmc_sec_title">Acknowledgments</h2>
    <p>Thanks.</p></section>
  <section id="notes1"><h2 class="pmc_sec_title">Author contributions</h2><p>All.</p></section>
  <section id="S2"><h2 class="pmc_sec_title">Supplementary Material</h2>
    <section class="sm xbox font-sm" id="SM1"><div class="media p"><div class="caption">
      <a href="/articles/instance/1/bin/s1.xlsx">s1.xlsx</a><sup> (12KB, xlsx) </sup>
      <p>Table S1. Y2H screen.</p></div></div></section>
  </section>
  <section class="ref-list" id="ref-list1"><h2 class="pmc_sec_title">References</h2>
    <ol><li><cite>A cited paper.</cite></li></ol></section>
  <section class="associated-data"><p>Data.</p></section>
</section></body></html>"""


def test_pmc_page(tmp_path: Path):
    """A PMC article page is read by its structure, as its XML: the same blocks, the back
    matter left out."""
    path = tmp_path / "a.html"
    path.write_text(PMC_PAGE)
    assert parse(path) == (
        "# NS1 binds TRIM25\n\n"
        "## Abstract\n\nWe show that NS1 binds TRIM25.\n\n"
        "## Results\n\n"
        "NS1 co-precipitates TRIM25 (Fig. 1).\n\n"
        "**Figure 1.** Co-IP. Anti-FLAG.\n\n"
        "**Table 1.** Interactions.\n\n"
        "| Bait | Prey | Prey |\n|---|---|---|\n| NS1 | TRIM25 | Q14258 |\n\n"
        "Co-IP in HEK293T cells.\n\n"
        "### Mapping\n\n- Residues 1-73 of NS1.\n\n"
        "## Supplementary Material\n\n"
        "**Supplementary material.** Table S1. Y2H screen. (file s1.xlsx)\n"
    )


def test_html_charset(tmp_path: Path):
    paragraphs = "".join(
        f"<p>Prot&eacute;ine {i} : NS1 lie TRIM25, voir la figure.</p>" for i in range(30)
    )
    head = '<head><meta charset="iso-8859-1"></head>'
    page = f"<html>{head}<body><article>{paragraphs}</article></body></html>"
    path = tmp_path / "a.html"
    path.write_bytes(page.replace("&eacute;", "\u00e9").encode("iso-8859-1"))
    assert "Prot\u00e9ine 29 : NS1 lie TRIM25" in parse(path)


def test_pdf_pages(tmp_path: Path):
    document = pymupdf.open()
    for words in ("NS1 binds TRIM25", "Figure 2: co-IP"):
        page = document.new_page()
        page.insert_text((72, 72), words)
    path = tmp_path / "a.pdf"
    document.save(path)
    assert parse(path) == (
        "--- page 1 ---\n\nNS1 binds TRIM25\n\n--- page 2 ---\n\nFigure 2: co-IP\n\n"
    )


def test_scanned_pdf(tmp_path: Path):
    document = pymupdf.open()
    document.new_page()
    path = tmp_path / "a.pdf"
    document.save(path)
    with pytest.raises(ParseError, match="no text layer"):
        parse(path)


def workbook(path: Path) -> None:
    """A workbook of two sheets, the second empty."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "Y2H; Caco-2"
    sheet.append(["Bait", "Prey", "Score", None])
    sheet.append(["NS1", "TRIM\u200b25", 2.0, None])
    sheet.append([None, None, None, None])
    sheet.append(["NS2", "line\nbreak", 0.5])
    book.create_sheet("Empty")
    book.save(path)


def test_xlsx(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    workbook(path)
    assert sheets(path) == [
        ("Y2H; Caco-2", "Bait\tPrey\tScore\nNS1\tTRIM25\t2\nNS2\tline break\t0.5\n"),
        ("Empty", ""),
    ]


def test_xls(tmp_path: Path):
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Y2H")
    for column, value in enumerate(["Bait", "Prey", "Score"]):
        sheet.write(0, column, value)
    sheet.write(1, 0, "NS1")
    sheet.write(
        1, 1, datetime(2002, 9, 2), xlwt.easyxf(num_format_str="D-MMM")
    )  # SEPT2, as Excel saw it
    sheet.write(1, 2, 2.0)
    path = tmp_path / "a.xls"
    workbook.save(str(path))
    assert sheets(path) == [("Y2H", "Bait\tPrey\tScore\nNS1\t2002-09-02T00:00:00\t2\n")]


def test_docx(tmp_path: Path):
    document = docx.Document()
    document.add_heading("Table S1", level=1)
    document.add_paragraph("Validated interactions.")
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, [("Bait", "Prey"), ("NS1", "TRIM25")], strict=True):
        for cell, value in zip(row.cells, cells, strict=True):
            cell.text = value
    path = tmp_path / "a.docx"
    document.save(str(path))
    assert parse(path) == (
        "# Table S1\n\nValidated interactions.\n\n| Bait | Prey |\n|---|---|\n| NS1 | TRIM25 |\n"
    )


def table_rows(markdown: str) -> list[list[str]]:
    """The rows of the Markdown tables of a converted text, their header separator left out."""
    lines = [line for line in markdown.splitlines() if line.startswith("| ")]
    return [line[2:-2].split(" | ") for line in lines]


def test_doc():
    """Table S3, Word 97-2003."""
    rows = table_rows(parse(SAMPLES / "ppat.1002331.s006.doc"))
    assert len(rows) == 3
    assert rows[0][1].startswith("Fraction of SARS-CoV interaction partners")
    assert rows[2] == [
        "SARS interaction partners (complete)",
        "36/514=7%",
        "884/31,941=2.8%",
        "4.4e-07",
    ]


def test_doc_rows_end_where_word_ends_them():
    """Table S1: a header cell broken over two lines, empty cells, and row properties too large
    for their page (in the Data stream)."""
    rows = table_rows(parse(SAMPLES / "ppat.1002331.s004.doc"))
    assert {len(r) for r in rows} == {9}
    assert rows[0][4] == "Number of times this prey has been isolated."
    assert rows[1] == [
        "A",
        "ORF13",
        "RYBP",
        "23429",
        "5",
        "RING1 and YY1 binding protein",
        "ORFeome",
        "4,9",
        "78,4",
    ]
    assert len(rows) == 88


def test_doc_cell_spanning_columns():
    """Table S2: empty last cells, and a row whose second cell spans two columns."""
    rows = table_rows(parse(SAMPLES / "ppat.1002331.s005.doc"))
    assert {len(r) for r in rows} == {7}
    assert rows[3] == ["Nsp13", "nsp13", "DDX5", "1655", "Ddx5", "19224332", ""]
    spanning = next(r for r in rows if "(prev desig U122 and X4)" in r)
    assert spanning == ["", "(prev desig U122 and X4)", "(prev desig U122 and X4)", "", "", "", ""]


def test_doc_characters_beyond_the_basic_plane():
    """A character beyond U+FFFF takes 4 bytes in the text: the cell marks after it are still
    found."""
    assert parse(SAMPLES / "beyond-bmp.doc") == (
        "Binding of \U0001d6fc-actinin (\U0001d6fc beyond U+FFFF) to NS1.\n\n"
        "| Bait | Prey |\n|---|---|\n| \U0001d6fc-NS1 | TRIM25 |\n\nAfter the table.\n"
    )


def test_doc_page_break_ends_a_paragraph():
    assert parse(SAMPLES / "page-break.doc") == (
        "Before the break.\n\nAfter the break.\n\nNext paragraph.\n"
    )


def test_not_a_doc(tmp_path: Path):
    path = tmp_path / "a.doc"
    path.write_text("plain text")
    with pytest.raises(ParseError, match="not a Word 97-2003 document"):
        parse(path)


def test_unknown_type(tmp_path: Path):
    path = tmp_path / "a.tif"
    path.write_bytes(b"II*\x00")
    with pytest.raises(ParseError, match="no converter for .tif"):
        parse(path)


def test_damaged_file_is_refused(tmp_path: Path):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"PK\x03\x04 not really a workbook")
    with pytest.raises(ParseError, match="cannot be read"):
        sheets(path)


# The texts, named by their format


def test_spreadsheet_gives_one_tsv_file_per_sheet(tmp_path: Path):
    """Numbered in workbook order, an empty sheet an empty file: every sheet listed."""
    (tmp_path / "1.supp").mkdir()
    workbook(tmp_path / "1.supp" / "s1.xlsx")
    rows = "Bait\tPrey\tScore\nNS1\tTRIM25\t2\nNS2\tline break\t0.5\n"
    assert convert(tmp_path, "1.supp/s1.xlsx") == {
        "file": "1.supp/s1.xlsx",
        "sheets": [
            {"sheet": "Y2H; Caco-2", "text": "1.supp/s1.xlsx.1.tsv", "bytes": len(rows)},
            {"sheet": "Empty", "text": "1.supp/s1.xlsx.2.tsv", "bytes": 0},
        ],
    }
    assert (tmp_path / "1.supp" / "s1.xlsx.1.tsv").read_text() == rows
    assert (tmp_path / "1.supp" / "s1.xlsx.2.tsv").read_text() == ""


def test_pdf_gives_plain_text_and_xml_markdown(tmp_path: Path):
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "NS1 binds TRIM25")
    document.save(tmp_path / "1.pdf")
    (tmp_path / "1.xml").write_text(JATS)
    pdf = convert(tmp_path, "1.pdf")
    assert pdf == {"file": "1.pdf", "text": "1.pdf.txt", "bytes": pdf["bytes"]}
    assert (tmp_path / "1.pdf.txt").read_text() == "--- page 1 ---\n\nNS1 binds TRIM25\n\n"
    xml = convert(tmp_path, "1.xml")
    assert xml == {"file": "1.xml", "text": "1.xml.md", "bytes": xml["bytes"]}
    assert xml["bytes"] == (tmp_path / "1.xml.md").stat().st_size


def test_text_files_are_their_own_text(tmp_path: Path):
    (tmp_path / "s1.csv").write_bytes(b"bait,prey\nNS1,TRIM25\n")
    assert convert(tmp_path, "s1.csv") == {"file": "s1.csv", "text": "s1.csv", "bytes": 21}
    assert [p.name for p in tmp_path.iterdir()] == ["s1.csv"]


@pytest.mark.parametrize(
    "raw",
    [
        "bait;prey\nNS1;prot\u00e9ine\n".encode("cp1252"),  # Excel's Windows export
        "bait;prey\nNS1;prot\u00e9ine\n".encode("utf-16"),  # Excel's "Unicode text"
        "bait;prey\nNS1;prot\u00e9ine\n".encode("mac_roman"),
        "\ufeffbait;prey\nNS1;prot\u200b\u00e9i\u00adne\n".encode(),
    ],
)
def test_text_files_are_cleaned_in_place(tmp_path: Path, raw: bytes):
    """In UTF-8, invisible characters removed."""
    (tmp_path / "s1.csv").write_bytes(raw)
    entry = convert(tmp_path, "s1.csv")
    assert (tmp_path / "s1.csv").read_text() == "bait;prey\nNS1;prot\u00e9ine\n"
    assert entry == {"file": "s1.csv", "text": "s1.csv", "bytes": 24}


def test_text_of_unknown_encoding_is_refused(tmp_path: Path):
    (tmp_path / "s1.txt").write_bytes(bytes(range(256)) * 4)
    with pytest.raises(ParseError, match="unknown encoding"):
        convert(tmp_path, "s1.txt")


def test_refused_file_writes_nothing(tmp_path: Path):
    (tmp_path / "s1.xlsx").write_bytes(b"PK\x03\x04 not really a workbook")
    with pytest.raises(ParseError, match="cannot be read"):
        convert(tmp_path, "s1.xlsx")
    with pytest.raises(ParseError, match="no such file"):
        convert(tmp_path, "s2.pdf")
    assert [p.name for p in tmp_path.iterdir()] == ["s1.xlsx"]


def test_command_prints_the_manifest_entries(tmp_path: Path):
    (tmp_path / "1.xml").write_text(JATS)
    (tmp_path / "s1.tif").write_bytes(b"II*\x00")
    workbook(tmp_path / "s2.xlsx")
    command = ["uv", "run", "drakkar-parse", str(tmp_path), "1.xml", "s1.tif", "s2.xlsx"]
    done = subprocess.run(command, capture_output=True, text=True)
    assert done.returncode == 1
    assert done.stderr == "s1.tif: not converted, no converter for .tif\n"
    entries = [json.loads(line) for line in done.stdout.splitlines()]
    assert [e.get("text") for e in entries] == ["1.xml.md", None]
    assert [s["sheet"] for s in entries[1]["sheets"]] == ["Y2H; Caco-2", "Empty"]
    assert (tmp_path / "1.xml.md").read_text().startswith("# NS1 binds TRIM25")
