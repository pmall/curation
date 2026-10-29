"""The derived datasets exported from the test database."""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from database.export import HH, INTERACTORS, MAPPINGS_H, MAPPINGS_V, VH, check_export, export
from database.generic import sync


def _write(path: Path, lines: list[str]) -> None:
    path.write_text("".join(line + "\n" for line in lines))


def _lines(path: Path) -> list[list[str]]:
    return [line.split("\t") for line in path.read_text().splitlines()]


def _generic(conn, directory: Path) -> Path:
    """The files of generic/ in step with the test database: a few species, some families as
    species so that every taxon is under one, each generic name its own protein."""
    directory.mkdir()
    _write(
        directory / "species.tsv",
        [
            "family_ncbi_taxon_id\tncbi_taxon_id\tname\tpublications",
            "\t694009\tSARS-CoV\t0",
            "\t3050299\tEBV\t0",
            "\t3052230\tHCV\t0",
            "\t2697049\tSARS-CoV-2\t0",
            "\t3048459\tZika virus\t0",
            "\t3430922\tMolluscum contagiosum virus\t0",
            "\t11118\tCoronaviridae\t0",
            "\t151340\tPapillomaviridae\t0",
            "\t11018\tTogaviridae\t0",
            "\t11270\tRhabdoviridae\t0",
            "\t11646\tLentivirus\t0",
        ],
    )
    _write(
        directory / "placed.tsv",
        ["ncbi_taxon_id\tspecies_ncbi_taxon_id\tname", "663565\t694009\tBat SARS Cov Rs806/2006"],
    )
    sync(conn.cursor(), directory)
    names = _lines(directory / "proteins.tsv")[1:]
    _write(
        directory / "proteins.tsv",
        ["ncbi_taxon_id\tgeneric_name\tprotein\tnote"]
        + [f"{t}\t{n}\t{n}\t" for t, n, _, _ in names],
    )
    assert sync(conn.cursor(), directory).endswith("To decide:\nnothing")
    return directory


def test_export(conn, tmp_path):
    """The datasets, named after the last vh run exported (ID 69, run 81), the mappings aside:
    the hh descriptions; the vh descriptions whose human interactor is up to date, with their
    species and generic viral protein; each of their viral interactors once, with its sequence.
    Each as `.tsv` and `.xlsx`, the same lines."""
    generic = _generic(conn, tmp_path / "generic")
    report = export(conn.cursor(), generic, tmp_path / "out")
    assert report.splitlines() == [
        "vh runs: 6, all; special runs: 0, all",
        "descriptions_hh_run81 (.tsv, .xlsx): 5 descriptions",
        "descriptions_vh_run81 (.tsv, .xlsx): 425 descriptions",
        "interactors_viral_run81 (.tsv, .xlsx): 171 viral interactors",
        "mappings_h_run81 (.tsv, .xlsx): 308 mappings",
        "mappings_v_run81 (.tsv, .xlsx): 424 mappings",
        "check:",
        "  vh descriptions: 427 live = 425 exported + 2 left out "
        "(a human interactor not up to date)",
        "  hh descriptions: 5 live = 5 exported + 0 left out (a human interactor not up to date)",
        "  mappings: 732 on the descriptions exported = 732 lines + 0 merged (the same sequence on "
        "both sides of a homodimer)",
        "  ok",
    ]

    vh = _lines(tmp_path / "out" / "descriptions_vh_run81.tsv")
    assert vh[0] == VH
    assert [r[0] for r in vh[1:]] == sorted(r[0] for r in vh[1:])
    # a strain placed in its species by placed.tsv
    assert [
        "EY08ABDB78", "2", "22046132", "MI:0018", "two hybrid", "2019_01", "Q56VL3", "OCIAD2",
        "694009", "SARS-CoV", "nsp13", "663565", "Bat SARS Cov Rs806/2006",
        "2019_01", "D2DJX6", "Helicase (Fragment)", "nsp13", "1", "601",
    ] in vh  # fmt: skip
    # a strain under a species of species.tsv, its deepest
    assert [
        "EY11D825B6", "1", "25031347", "MI:0004", "affinity chromatography technology",
        "2019_01", "P55209", "NAP1L1", "3052230", "HCV", "NS5A", "356411",
        "Hepatitis C virus JFH-1", "2019_01", "Q99IB8", "Genome polyprotein", "NS5A",
        "1977", "2442",
    ] in vh  # fmt: skip
    cur = conn.cursor()
    cur.execute(
        """
        SELECT d.stable_id
        FROM descriptions AS d
        JOIN proteins AS h ON h.id = d.protein1_id
        LEFT JOIN proteins_versions AS current1
          ON current1.accession = h.accession AND current1.version = h.version
        WHERE d.deleted_at IS NULL AND current1.version IS NULL
        """
    )
    obsolete = {stable_id for (stable_id,) in cur.fetchall()}
    assert obsolete and not obsolete & {r[0] for r in vh}  # a human interactor not up to date

    interactors = _lines(tmp_path / "out" / "interactors_viral_run81.tsv")
    assert interactors[0] == INTERACTORS
    keys = [tuple(r[:4]) for r in interactors[1:]]
    assert len(keys) == len(set(keys)) == len({(r[13], r[14], r[17], r[18]) for r in vh[1:]})
    assert set(keys) == {(r[13], r[14], r[17], r[18]) for r in vh[1:]}
    cur.execute(
        """
        SELECT substr(sequences ->> accession, 1977, 2442 - 1977 + 1)
        FROM proteins
        WHERE accession = 'Q99IB8' AND version = '2019_01'
        """
    )
    (ns5a,) = cur.fetchone()
    assert ["2019_01", "Q99IB8", "1977", "2442", ns5a] in interactors  # a mature protein

    hh = _lines(tmp_path / "out" / "descriptions_hh_run81.tsv")
    assert hh[0] == HH
    assert hh[1] == [
        "EY1F03BBEB", "2", "15471871", "MI:0515", "methyltransferase assay",
        "2019_01", "Q96RI1", "NR1H4", "2020_03", "P68431", "H3C1",
    ]  # fmt: skip

    # the .xlsx: the same lines, numbers as numbers
    for name, lines in [
        ("descriptions_hh_run81", hh),
        ("descriptions_vh_run81", vh),
        ("interactors_viral_run81", interactors),
    ]:
        sheet = load_workbook(tmp_path / "out" / f"{name}.xlsx", read_only=True)[name]
        rows = list(sheet.iter_rows(values_only=True))
        assert [["" if v is None else str(v) for v in row] for row in rows] == lines
    assert rows[1][2] == 1 and isinstance(rows[1][2], int)  # interactors_viral_run81: start


def test_runs(conn, tmp_path):
    """`run` keeps the vh runs up to its ID, special runs aside, `special_run` the special runs
    up to its ID, all by default; hh is whole. The files are named after the last vh run, its
    name sanitized. A cut that is not a run of its kind is refused."""
    generic = _generic(conn, tmp_path / "generic")
    cur = conn.cursor()
    cur.execute("UPDATE runs SET name = '32 b/2' WHERE id = 12")  # a name to sanitize
    cur.execute("INSERT INTO runs (id, type, name) VALUES (70, 'vh', 'special-1')")
    cur.execute("INSERT INTO runs (id, type, name) VALUES (71, 'vh', 'special-2')")
    cur.execute("UPDATE associations SET run_id = 70 WHERE pmid = 24379396")  # from run ID 30
    cur.execute("UPDATE associations SET run_id = 71 WHERE pmid = 31690057")  # from run ID 44

    def pmids(path: Path) -> set[str]:
        return {r[2] for r in _lines(path)[1:]}

    report = export(cur, generic, tmp_path / "a", run=12)
    assert report.splitlines()[0] == "vh runs: 2, up to ID 12; special runs: 2, all"
    assert pmids(tmp_path / "a" / "descriptions_vh_run32-b-2.tsv") == {
        "8626429", "9858596", "22046132", "25031347",  # run IDs 1 and 12; not 10777214, its
        # human interactors not up to date
        "24379396", "31690057",  # the special runs
    }  # fmt: skip
    assert pmids(tmp_path / "a" / "descriptions_hh_run32-b-2.tsv") == {"15471871", "7499434"}

    report = export(cur, generic, tmp_path / "b", run=12, special_run=70)
    assert report.splitlines()[0] == "vh runs: 2, up to ID 12; special runs: 1, up to ID 70"
    assert "31690057" not in pmids(tmp_path / "b" / "descriptions_vh_run32-b-2.tsv")
    assert "24379396" in pmids(tmp_path / "b" / "descriptions_vh_run32-b-2.tsv")

    for run, special_run, message in [
        (31, None, "--run 31: not the ID of a vh run"),  # an hh run
        (70, None, "--run 70: not the ID of a vh run"),  # a special run
        (None, 12, "--special-run 12: not the ID of a special run"),
    ]:
        with pytest.raises(SystemExit, match=message):
            export(cur, generic, tmp_path / "c", run=run, special_run=special_run)


def test_not_in_step(conn, tmp_path):
    """Nothing is written when generic/ misses the species of a taxon or the protein of a name."""
    generic = _generic(conn, tmp_path / "generic")
    species = (generic / "species.tsv").read_text()
    (generic / "species.tsv").write_text(species.replace("\t11646\tLentivirus\t1\n", ""))
    proteins = (generic / "proteins.tsv").read_text()
    proteins = proteins.replace("3052230\tNS5A\tNS5A", "3052230\tNS5A\t")
    (generic / "proteins.tsv").write_text(proteins)
    with pytest.raises(SystemExit) as refused:
        export(conn.cursor(), generic, tmp_path / "out")
    lines = str(refused.value).splitlines()
    assert "taxon under no species\t11706" in lines  # HIV-1 HXB2R, Lentivirus gone
    assert "no protein\t3052230 HCV\tNS5A" in lines
    assert not (tmp_path / "out").exists()


def _occurrence(accession: str) -> dict:
    return {"accession": accession, "occurrences": [{"start": 1, "stop": 4, "identity": 100}]}


def test_mappings(conn, tmp_path):
    """Each mapping with its interactor and the isoform it is found on: the canonical one when the
    mapping occurs on it, else the other isoform with the lowest number. One line per (stable ID,
    sequence, isoform): a homodimer's mapping on both sides is one line. `target`, in
    `mappings_h`, is the kind of the other interactor."""
    generic = _generic(conn, tmp_path / "generic")
    cur = conn.cursor()
    # an hh description made a homodimer of NR1H4, the same mapping on both sides
    homodimer = [{"sequence": "MSSN", "isoforms": [_occurrence("Q96RI1")]}]
    cur.execute(
        """
        UPDATE descriptions SET protein2_id = protein1_id, mapping1 = %s, mapping2 = %s
        WHERE stable_id = 'EYD107712D' AND deleted_at IS NULL
        """,
        (json.dumps(homodimer), json.dumps(homodimer)),
    )
    # NAP1L1 × HCV NS5A: a human mapping only on isoforms, a viral one with extra end residues,
    # in lowercase in the export
    human = [
        {"sequence": "MADI", "isoforms": [_occurrence("P55209")]},
        {"sequence": "AAAA", "isoforms": [_occurrence("P55209-10"), _occurrence("P55209-2")]},
    ]
    viral = [{"sequence": "MSGSWK", "nter": "M", "cter": "K", "isoforms": [_occurrence("Q99IB8")]}]
    cur.execute(
        """
        UPDATE descriptions SET mapping1 = %s, mapping2 = %s
        WHERE stable_id = 'EY11D825B6' AND deleted_at IS NULL
        """,
        (json.dumps(human), json.dumps(viral)),
    )
    report = export(cur, generic, tmp_path / "out").splitlines()

    h = _lines(tmp_path / "out" / "mappings_h_run81.tsv")
    v = _lines(tmp_path / "out" / "mappings_v_run81.tsv")
    assert h[0] == MAPPINGS_H and v[0] == MAPPINGS_V
    assert h[1:] == sorted(h[1:]) and v[1:] == sorted(v[1:])
    assert [line for line in h if line[0] == "EYD107712D"] == [
        ["EYD107712D", "MSSN", "2019_01", "Q96RI1", "Q96RI1", "h", "", ""]
    ]
    assert [line for line in h if line[0] == "EY11D825B6"] == [
        ["EY11D825B6", "AAAA", "2019_01", "P55209", "P55209-2", "v", "", ""],
        ["EY11D825B6", "MADI", "2019_01", "P55209", "P55209", "v", "", ""],
    ]
    assert [line for line in v if line[0] == "EY11D825B6"] == [
        ["EY11D825B6", "mSGSWk", "2019_01", "Q99IB8", "1977", "2442", "Q99IB8", "M", "K"]
    ]
    descriptions = {
        r[0]
        for name in ("descriptions_hh_run81", "descriptions_vh_run81")
        for r in _lines(tmp_path / "out" / f"{name}.tsv")[1:]
    }
    assert {line[0] for line in h[1:] + v[1:]} <= descriptions
    assert f"mappings_h_run81 (.tsv, .xlsx): {len(h) - 1} mappings" in report


def test_check(conn, tmp_path):
    """The files of an export, read back, are checked against each other and against the
    database: each corruption is a failure."""
    generic = _generic(conn, tmp_path / "generic")
    out = tmp_path / "out"
    export(conn.cursor(), generic, out)

    def edit(name: str, change) -> None:
        path = out / f"{name}_run81.tsv"
        lines = path.read_text().splitlines()
        path.write_text("".join(line + "\n" for line in change(lines)))

    edit("descriptions_hh", lambda lines: [*lines, lines[1]])  # a description repeated
    edit("interactors_viral", lambda lines: [lines[0], *lines[2:]])  # an interactor gone
    first = _lines(out / "interactors_viral_run81.tsv")[1]
    edit(  # a sequence cut
        "interactors_viral",
        lambda lines: [
            line[:-1] if line.startswith("\t".join(first[:4])) else line for line in lines
        ],
    )
    edit("mappings_v", lambda lines: [*lines, lines[1]])  # a mapping repeated
    mapping = _lines(out / "mappings_h_run81.tsv")[1]
    edit(  # a mapping on another entry
        "mappings_h",
        lambda lines: [line.replace(f"\t{mapping[3]}\t", "\tP99999\t", 1) for line in lines],
    )

    numbers, failures = check_export(conn.cursor(), out, "run81", [1, 12, 30, 44, 56, 69])
    assert (
        "hh descriptions: 5 live ≠ 6 exported + 0 left out (a human interactor not up to date)"
        in numbers
    )
    assert f"description on 2 lines: {_lines(out / 'descriptions_hh_run81.tsv')[1][0]}" in failures
    assert f"sequence not as long as its coordinates: {' '.join(first[:4])}" in failures
    assert any(f.startswith("viral interactor without its line: ") for f in failures)
    assert any(f.startswith("mappings_v: mapping on several lines: ") for f in failures)
    assert f"mappings_h: {mapping[2]} P99999 not an interactor of {mapping[0]}" in failures
    assert f"mappings_h: isoform {mapping[4]} not of P99999" in failures
    assert "hh descriptions: the numbers do not add up" in failures
    assert "mappings: the numbers do not add up" in failures
