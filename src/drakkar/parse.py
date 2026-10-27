"""Convert full texts and supplementary files to the texts a curation reads, short in tokens.

uv run drakkar-parse <dir> <file>...

PMC XML, HTML (a PMC article page by its structure, as its XML) and Word documents (.docx,
and .doc through our own reader) become Markdown, `<file>.md`; a PDF its plain text page by
page, a table's row on one line when the PDF writes it cell after cell, `<file>.txt`; a
spreadsheet one TSV file per sheet, `<file>.<n>.tsv` in workbook order, every row as wide as
its widest.
CSV, TSV and plain text are their own text, cleaned in place: in UTF-8, invisible characters
removed, a CSV's or TSV's rows padded as a sheet's. The texts are written beside the file,
named by their format; the converter is chosen by the file type; no network. Prints each
file's entry of the manifest (`drakkar-fetch`), its paths relative to `<dir>`.
"""

import argparse
import bisect
import csv
import io
import json
import re
import struct
import sys
from collections import Counter
from collections.abc import Callable, Collection, Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Literal, cast

import charset_normalizer
import docx
import olefile
import openpyxl
import pymupdf
import trafilatura
import xlrd
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree
from lxml.etree import _Element  # pyright: ignore[reportPrivateUsage]


class ParseError(Exception):
    pass


def one_line(text: str) -> str:
    return " ".join(text.split())


def markdown_table(rows: list[list[str]]) -> str:
    """Rows of cells as a Markdown table, the first row as its header."""
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    cells = [[one_line(c).replace("|", "\\|") for c in r] + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(r) + " |" for r in cells]
    lines.insert(1, "|" + "---|" * width)
    return "\n".join(lines)


# PMC XML (JATS)

# Elements that are blocks of their own, even inside a paragraph.
BLOCKS = {
    "sec",
    "p",
    "list",
    "fig",
    "table-wrap",
    "supplementary-material",
    "boxed-text",
    "disp-formula",
    "def-list",
    "disp-quote",
    "statement",
    "fig-group",
    "table-wrap-group",
}
# Elements whose text is set apart from the text around it, as a block's is.
SEPARATE = BLOCKS | {"term", "def", "def-item", "list-item", "fn", "td", "th", "label", "title"}
# Never read for a curation.
DROPPED = {
    "ref-list",
    "ack",
    "fn-group",
    "notes",
    "sub-article",
    "response",
    "object-id",
    "permissions",
}
XLINK = "{http://www.w3.org/1999/xlink}href"


def tag_name(el: _Element) -> str:
    # Comments and processing instructions have a function as tag.
    return etree.QName(el).localname if isinstance(el.tag, str) else ""  # pyright: ignore[reportUnnecessaryIsInstance]


def _text(el: _Element, skip: Collection[str] = ()) -> str:
    """All the text of an element in one line, but for its children whose tag is in `skip`."""
    parts = [el.text or ""]
    for child in el:
        tag = tag_name(child)
        if tag == "alternatives":  # one formula in several notations: the first
            notations = [c for c in child if tag_name(c)]
            parts.append(_text(notations[0], skip) if notations else "")
        elif tag in SEPARATE and tag not in skip:
            parts.append(f" {_text(child, skip)} ")
        elif tag not in skip:
            parts.append(_text(child, skip))
        parts.append(child.tail or "")
    return one_line("".join(parts))


def _inline(el: _Element) -> str:
    """The text of an element, without its nested blocks."""
    return _text(el, BLOCKS)


def _caption(el: _Element, kind: str) -> str:
    label = el.find("{*}label")
    caption = el.find("{*}caption")
    head = _inline(label) if label is not None else kind
    text = ""
    if caption is not None:
        text = " ".join(_text(c) for c in caption if tag_name(c) in ("title", "p"))
    return f"**{head}.** {text}".rstrip()


def _span(value: str | None) -> int:
    try:
        return max(int(value or 1), 1)
    except ValueError:
        return 1


def _grid(table: _Element) -> list[list[str]]:
    """The cells of an XHTML table, spans spread so that every row has its columns."""
    rows: list[list[str]] = []
    pending: dict[tuple[int, int], str] = {}  # cells of rowspans, by (row, column)
    for r, tr in enumerate(table.iter("{*}tr")):
        row: list[str] = []
        cells = iter([c for c in tr if tag_name(c) in ("td", "th")])
        cell = next(cells, None)
        while cell is not None or (r, len(row)) in pending:
            if (r, len(row)) in pending:
                row.append(pending.pop((r, len(row))))
                continue
            assert cell is not None
            text = _text(cell)
            colspan, rowspan = _span(cell.get("colspan")), _span(cell.get("rowspan"))
            for c in range(len(row), len(row) + colspan):
                for below in range(r + 1, r + rowspan):
                    pending[(below, c)] = text
            row.extend([text] * colspan)
            cell = next(cells, None)
        rows.append(row)
    return rows


def _tables(tables: Iterable[_Element]) -> list[str]:
    """The tables in Markdown; without any, the table is given as an image."""
    return [markdown_table(_grid(table)) for table in tables] or ["(table given as an image)"]


def _jats(children: Iterable[_Element], level: int) -> Iterator[str]:
    for child in children:
        tag = tag_name(child)
        if tag in DROPPED or tag in ("label", "title", "caption"):
            continue
        if tag in ("sec", "app", "abstract", "trans-abstract", "boxed-text", "app-group"):
            title = child.find("{*}title")
            if title is not None and _inline(title):
                yield "#" * min(level, 6) + " " + _inline(title)
            elif tag in ("abstract", "trans-abstract"):
                yield "#" * min(level, 6) + " Abstract"
            yield from _jats(child, level + 1)
        elif tag == "fig":
            yield _caption(child, "Figure")
        elif tag == "table-wrap":
            yield _caption(child, "Table")
            yield from _tables(child.iter("{*}table"))
            for foot in child.iter("{*}table-wrap-foot"):
                yield _text(foot)
        elif tag == "supplementary-material":
            media = child.find(".//{*}media")
            href = child.get(XLINK) or (media.get(XLINK) if media is not None else None)
            yield _caption(child, "Supplementary material") + (f" (file {href})" if href else "")
        elif tag == "list":
            items = child.iter("{*}list-item")
            yield "\n".join(f"- {_text(item, {'list'})}" for item in items)  # nested: own items
        elif tag == "def-list":
            yield "\n".join(f"- {_text(item)}" for item in child.iter("{*}def-item"))
        elif tag == "alternatives":  # one block in several notations: the first
            yield from _jats([c for c in child if tag_name(c)][:1], level)
        elif blocks := [c for c in child if tag_name(c) in BLOCKS]:
            if text := _inline(child):
                yield text
            yield from _jats(blocks, level)
        elif text := _inline(child):
            yield text


def jats(path: Path) -> str:
    try:
        root = etree.parse(path, etree.XMLParser(resolve_entities=False, no_network=True))
    except etree.XMLSyntaxError as e:
        raise ParseError(f"not XML: {e}") from e
    article = root.getroot()
    if tag_name(article) != "article":
        article = article.find(".//{*}article")
        if article is None:
            raise ParseError("XML without an article")
    blocks: list[str] = []
    title = article.find("{*}front/{*}article-meta/{*}title-group/{*}article-title")
    if title is not None:
        blocks.append("# " + _inline(title))
    meta = article.find("{*}front/{*}article-meta")
    if meta is not None:
        blocks.extend(_jats([a for a in meta if tag_name(a) in ("abstract", "trans-abstract")], 2))
    for part in ("body", "back", "floats-group"):
        el = article.find(f"{{*}}{part}")
        if el is not None:
            blocks.extend(_jats(el, 2))
    return "\n\n".join(b for b in blocks if b) + "\n"


# HTML

# The back matter of a PMC article page, left out as in its XML; `notes-<n>` sections too.
PMC_PAGE_DROPPED = {"kwd-group", "ack", "fn-group", "ref-list", "associated-data", "contrib-info"}


def class_names(el: _Element) -> set[str]:
    return set((el.get("class") or "").split())


def _pmc_page(children: Iterable[_Element], level: int) -> Iterator[str]:
    """The blocks of a PMC article page, its XML rendered as HTML: sections and their titles,
    paragraphs, lists, figure and table captions, tables, supplementary materials."""
    for child in children:
        tag, classes = tag_name(child), class_names(child)
        if tag == "section":
            if classes & PMC_PAGE_DROPPED or (child.get("id") or "").startswith("notes"):
                continue
            if "tw" in classes:  # a table
                head = child.find("h4")
                caption = child.find("div[@class='caption p']")
                label = _text(head) if head is not None else "Table."
                yield f"**{label}** {_text(caption) if caption is not None else ''}".rstrip()
                yield from _tables(child.iter("table"))
                for foot in child.iterfind(".//div"):
                    if "tw-foot" in class_names(foot):
                        yield _text(foot)
            elif "sm" in classes:  # a supplementary material: its caption and file
                links = [a.get("href", "") for a in child.iter("a")]
                names = list(dict.fromkeys(Path(h).name for h in links if "/bin/" in h))
                caption = _text(child, {"a", "sup"})
                files = f"(file {', '.join(names)})" if names else ""
                yield " ".join(x for x in ("**Supplementary material.**", caption, files) if x)
            else:
                title = next((h for h in child if tag_name(h) in ("h2", "h3", "h4", "h5")), None)
                if title is not None and _text(title):
                    yield "#" * min(level, 6) + " " + _text(title)
                yield from _pmc_page(child, level + 1)
        elif tag == "figure":
            head = next((h for h in child if tag_name(h) in ("h3", "h4")), None)
            caption = child.find("figcaption")
            label = _text(head) if head is not None else "Figure."
            yield f"**{label}** {_text(caption) if caption is not None else ''}".rstrip()
        elif tag in ("h2", "h3", "h4", "h5"):
            continue  # a section title, given with its section
        elif tag in ("ul", "ol"):
            yield "\n".join(f"- {_text(item)}" for item in child.iterfind("li"))
        elif tag == "table":
            yield markdown_table(_grid(child))
        elif tag in ("div", "dl") and any(
            tag_name(c) in ("p", "section", "figure", "table", "ul", "ol", "div") for c in child
        ):
            yield from _pmc_page(child, level)
        elif text := _text(child):
            yield text


def pmc_page(doc: _Element, body: _Element) -> str:
    title = next(doc.iter("h1"), None)
    blocks = ["# " + _text(title)] if title is not None else []
    blocks.extend(_pmc_page(body, 2))
    return "\n\n".join(b for b in blocks if b) + "\n"


def html(path: Path) -> str:
    """A PMC article page by its structure, as its XML; any other page by its main text."""
    raw = path.read_bytes()
    if b"main-article-body" in raw:
        doc = etree.fromstring(raw, etree.HTMLParser())
        for body in doc.iter("section"):
            if "main-article-body" in class_names(body):
                return pmc_page(doc, body)
    text = trafilatura.extract(
        raw,  # bytes: trafilatura finds the page's charset
        output_format="markdown",
        include_tables=True,
        include_comments=False,
        include_images=False,
        include_links=False,
        favor_recall=True,
    )
    if not text:
        raise ParseError("no main text in the HTML page")
    return text + "\n"


# PDF


@dataclass
class _Line:
    """A line of a page as MuPDF reads it: its text, the baseline and the size of its largest
    characters, and whether it is horizontal (not rotated)."""

    text: str
    baseline: float
    size: float
    horizontal: bool


def _beside(a: _Line, b: _Line) -> bool:
    """On one line of the page: both horizontal, their baselines closer than half the smaller
    size."""
    return a.horizontal and b.horizontal and abs(a.baseline - b.baseline) < min(a.size, b.size) / 2


def _page(page: Any) -> str:
    """A page's text as MuPDF reads it, line by line, in its order; a line read right after
    another and beside it (the next cell of a table's row) continues it, after a tab. Nothing
    is moved: two columns of text are never mixed."""
    rows: list[str] = []
    before: _Line | None = None
    blocks = page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]  # pyright: ignore[reportUnknownMemberType]
    for line in (line for block in blocks for line in block["lines"]):
        spans = [s for s in line["spans"] if s["text"].strip()]
        if not spans:
            continue
        largest = max(spans, key=lambda s: s["size"])
        text = "".join(s["text"] for s in line["spans"]).strip()
        current = _Line(text, largest["origin"][1], largest["size"], line["dir"][0] > 0.99)
        if before is not None and _beside(before, current):
            rows[-1] += "\t" + text
        else:
            rows.append(text)
        before = current
    return "\n".join(rows)


def pdf(path: Path) -> str:
    try:
        document = pymupdf.open(path)
    except pymupdf.FileDataError as e:
        raise ParseError(f"not a PDF: {e}") from e
    with document:
        pages = [_page(page) for page in document]
    if not any(pages):
        raise ParseError("no text layer: a scanned PDF")
    return "".join(f"--- page {n} ---\n\n{p}\n\n" for n, p in enumerate(pages, 1))


# Spreadsheets


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date):
        return value.isoformat()
    return one_line(str(value))


def _padded(rows: Iterable[list[str]]) -> list[list[str]]:
    """A table's rows, each as many cells as its widest row: its empty rows, and the empty
    cells after the last that holds a value, left out."""
    table: list[list[str]] = []
    for cells in rows:
        while cells and not cells[-1].strip():
            cells = cells[:-1]
        if cells:
            table.append(cells)
    width = max((len(cells) for cells in table), default=0)
    return [cells + [""] * (width - len(cells)) for cells in table]


def _tsv(rows: list[list[Any]]) -> str:
    """A sheet's rows, tab-separated, padded."""
    return "".join(
        "\t".join(cells) + "\n" for cells in _padded([_cell(v) for v in r] for r in rows)
    )


def xlsx(path: Path) -> list[tuple[str, str]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            (ws.title, _tsv([list(r) for r in ws.iter_rows(values_only=True)]))
            for ws in workbook.worksheets
        ]
    finally:
        workbook.close()


def _xls_value(cell: Any, datemode: Literal[0, 1]) -> Any:
    """A cell's value; a date as a date, as in .xlsx (a number in the file)."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, datemode)
        except xlrd.xldate.XLDateError:
            pass
    return cell.value


def xls(path: Path) -> list[tuple[str, str]]:
    book: Any = xlrd.open_workbook(str(path))
    sheets: list[tuple[str, str]] = []
    for sheet in book.sheets():
        rows = [[_xls_value(c, book.datemode) for c in sheet.row(i)] for i in range(sheet.nrows)]
        sheets.append((str(sheet.name), _tsv(rows)))
    return sheets


# Word


def word(path: Path) -> str:
    document = docx.Document(str(path))
    blocks: list[str] = []
    body = cast(_Element, document.element.body)  # pyright: ignore[reportAttributeAccessIssue, reportUnknownMemberType]
    for el in body.iterchildren():
        tag = etree.QName(el).localname
        if tag == "p":
            paragraph = Paragraph(cast(Any, el), document)
            text = one_line(paragraph.text)
            style = paragraph.style.name if paragraph.style is not None else ""
            if text and (m := re.match(r"Heading (\d)", style or "")):
                blocks.append("#" * int(m.group(1)) + " " + text)
            elif text:
                blocks.append(text)
        elif tag == "tbl":
            rows = [[c.text for c in row.cells] for row in Table(cast(Any, el), document).rows]
            blocks.append(markdown_table(rows))
    return "\n\n".join(blocks) + "\n"


# Old Word documents (Word 97-2003, [MS-DOC])

# Paragraph properties that place a paragraph in a table.
SPRM_IN_TABLE, SPRM_ROW_END, SPRM_DEPTH = 0x2416, 0x2417, 0x6649
SPRM_TABLE_DEFINITIONS = (0xD608, 0xD606)  # sprmTDefTable, sprmTDefTable10
SPRM_TABS = 0xC615  # sprmPChgTabs
SPRM_HUGE_PAPX = 0x6646  # the properties are in the Data stream, at this offset


@dataclass
class _Piece:
    """A run of the main text stored in one place of the WordDocument stream."""

    cp: int  # its first character
    length: int
    fc: int  # the byte offset of its first character
    width: int  # bytes per character: 1 (cp1252) or 2 (UTF-16)


def _u16(data: bytes, offset: int) -> int:
    return struct.unpack_from("<H", data, offset)[0]


def _u32(data: bytes, offset: int) -> int:
    return struct.unpack_from("<I", data, offset)[0]


@dataclass
class _Properties:
    """A paragraph's properties: the values of 1 to 4 bytes, and for the end of a table row the
    edges of its cells (sprmTDefTable)."""

    values: dict[int, int]
    edges: tuple[int, ...] = ()


def _sprms(grpprl: bytes) -> _Properties:
    """The properties set by a list of property modifiers."""
    found = _Properties({})
    i = 0
    while i + 2 <= len(grpprl):
        sprm = _u16(grpprl, i)
        i += 2
        spra = sprm >> 13
        if spra in (0, 1):
            size = 1
        elif spra in (2, 4, 5):
            size = 2
        elif spra == 3:
            size = 4
        elif spra == 7:
            size = 3
        elif sprm in SPRM_TABLE_DEFINITIONS:
            size = 1 + _u16(grpprl, i)  # 2 bytes of size, counting itself minus 1
        elif sprm == SPRM_TABS and i < len(grpprl) and grpprl[i] == 255:
            deleted = grpprl[i + 1]  # tabs deleted (positions, closures), then tabs added
            added = grpprl[i + 2 + 4 * deleted]
            size = 3 + 4 * deleted + 3 * added
        else:
            size = 1 + grpprl[i] if i < len(grpprl) else 0
        if size in (1, 2, 4) and i + size <= len(grpprl):
            found.values[sprm] = int.from_bytes(grpprl[i : i + size], "little")
        elif sprm in SPRM_TABLE_DEFINITIONS and i + size <= len(grpprl):
            cells = grpprl[i + 2]  # then the edges of the cells, in twips
            found.edges = struct.unpack_from(f"<{cells + 1}h", grpprl, i + 3)
        i += size
    return found


def _paragraphs(
    document: bytes, table: bytes, data: bytes, fc_lcb: int
) -> list[tuple[int, int, _Properties]]:
    """Every paragraph's byte range in the WordDocument stream and its properties, from the
    formatted disk pages listed in PlcBtePapx; properties too large for their page are in the
    Data stream (sprmPHugePapx)."""
    fc_plc, lcb_plc = struct.unpack_from("<II", document, fc_lcb + 13 * 8)
    plc = table[fc_plc : fc_plc + lcb_plc]
    n = (lcb_plc - 4) // 8
    paragraphs: list[tuple[int, int, _Properties]] = []
    for k in range(n):
        page = (_u32(plc, 4 * (n + 1) + 4 * k) & 0x3FFFFF) * 512
        fkp = document[page : page + 512]
        if len(fkp) < 512:
            raise ParseError("a Word document with a broken paragraph table")
        runs = fkp[511]
        fcs = struct.unpack_from(f"<{runs + 1}I", fkp, 0)
        for r in range(runs):
            offset = 2 * fkp[4 * (runs + 1) + 13 * r]
            properties = _Properties({})
            if offset:
                cb = fkp[offset]
                start, size = (offset + 1, 2 * cb - 1) if cb else (offset + 2, 2 * fkp[offset + 1])
                properties = _sprms(fkp[start + 2 : start + size])  # after the style index
                if (huge := properties.values.get(SPRM_HUGE_PAPX)) is not None:
                    if huge + 2 > len(data):
                        raise ParseError("a Word document with its properties out of its file")
                    properties = _sprms(data[huge + 2 : huge + 2 + _u16(data, huge)])
            paragraphs.append((fcs[r], fcs[r + 1], properties))
    paragraphs.sort(key=lambda p: p[0])
    return paragraphs


def _word_document(path: Path) -> tuple[str, list[int], list[tuple[int, int, _Properties]]]:
    """The main text of a Word 97-2003 document, the byte offset of each of its characters, and
    its paragraphs."""
    try:
        ole = olefile.OleFileIO(str(path))
    except OSError as e:
        raise ParseError(f"not a Word 97-2003 document: {e}") from e
    with ole:
        if not ole.exists("WordDocument"):
            raise ParseError("not a Word 97-2003 document")
        document = ole.openstream("WordDocument").read()
        if _u16(document, 0) != 0xA5EC or _u16(document, 2) < 0x00C1:
            raise ParseError("a Word document older than Word 97")
        flags = _u16(document, 0x0A)
        if flags & 0x0100:
            raise ParseError("an encrypted Word document")
        table_name = "1Table" if flags & 0x0200 else "0Table"
        if not ole.exists(table_name):
            raise ParseError("a Word document without its table stream")
        table = ole.openstream(table_name).read()
        data = ole.openstream("Data").read() if ole.exists("Data") else b""
    # The FIB: 32 bytes, then 16-bit values, 32-bit values, then (fc, lcb) pairs.
    fc_lw = 34 + 2 * _u16(document, 32) + 2
    fc_lcb = fc_lw + 4 * _u16(document, fc_lw - 2) + 2
    ccp_text = _u32(document, fc_lw + 12)
    fc_clx, lcb_clx = struct.unpack_from("<II", document, fc_lcb + 33 * 8)
    clx = table[fc_clx : fc_clx + lcb_clx]
    i = 0
    while i < len(clx) and clx[i] == 0x01:  # property modifiers, before the piece table
        i += 3 + _u16(clx, i + 1)
    if i >= len(clx) or clx[i] != 0x02:
        raise ParseError("a Word document without a piece table")
    lcb = _u32(clx, i + 1)
    plc = clx[i + 5 : i + 5 + lcb]
    n = (lcb - 4) // 12
    cps = struct.unpack_from(f"<{n + 1}I", plc, 0)
    pieces: list[_Piece] = []
    for k in range(n):
        fc = _u32(plc, 4 * (n + 1) + 8 * k + 2)
        compressed = bool(fc & 0x40000000)
        start = (fc & 0x3FFFFFFF) // 2 if compressed else fc
        pieces.append(_Piece(cps[k], cps[k + 1] - cps[k], start, 1 if compressed else 2))
    chars: list[str] = []
    offsets: list[int] = []
    for piece in pieces:
        if piece.cp >= ccp_text:
            break
        length = min(piece.length, ccp_text - piece.cp)
        raw = document[piece.fc : piece.fc + piece.width * length]
        decoded = raw.decode("cp1252" if piece.width == 1 else "utf-16-le", errors="replace")
        chars.append(decoded)
        position = piece.fc
        for char in decoded:  # a character beyond U+FFFF takes 4 bytes in UTF-16
            offsets.append(position)
            position += 1 if piece.width == 1 else len(char.encode("utf-16-le"))
    return "".join(chars), offsets, _paragraphs(document, table, data, fc_lcb)


def _aligned(rows: list[tuple[list[str], tuple[int, ...]]]) -> list[list[str]]:
    """The rows of one table, aligned. A cell spanning several columns is one wider cell in
    Word: in a row with fewer cells than the fullest row, each cell is placed by its edges on the
    nearest edges of the fullest row, its value repeated over the columns it spans."""
    grid = max((edges for _, edges in rows), key=len, default=())
    aligned: list[list[str]] = []
    for cells, edges in rows:
        if len(cells) < len(grid) - 1 and len(edges) == len(cells) + 1:
            column = [min(range(len(grid)), key=lambda k: abs(grid[k] - e)) for e in edges]
            cells = [
                c
                for c, a, b in zip(cells, column, column[1:], strict=False)
                for _ in range(max(b - a, 1))
            ]
        aligned.append(cells)
    return aligned


def doc(path: Path) -> str:
    """An old Word document, as Markdown like a .docx: its paragraphs, and its tables. Each mark
    that ends a paragraph (a carriage return, a cell mark in a table, a page or section break)
    is read with its paragraph's properties: in a table or not, the end of a cell or of a row. A
    nested table's cells are flattened into the cell holding them."""
    text, offsets, paragraphs = _word_document(path)
    starts = [p[0] for p in paragraphs]

    def properties(i: int) -> _Properties:
        k = bisect.bisect_right(starts, offsets[i]) - 1
        if k >= 0 and paragraphs[k][0] <= offsets[i] < paragraphs[k][1]:
            return paragraphs[k][2]
        return _Properties({})

    blocks: list[str] = []
    table: list[tuple[list[str], tuple[int, ...]]] = []  # the rows of the current table
    row: list[str] = []
    current: list[str] = []

    def end_paragraph() -> None:
        """The table before the paragraph, if any, then the paragraph."""
        nonlocal table, row, current
        if row:  # cells whose row end was not found: kept, never lost
            table.append((row, ()))
        blocks.append(markdown_table(_aligned(table)))
        blocks.append(one_line("".join(current)))
        table, row, current = [], [], []

    fields: list[bool] = []  # open fields, True while in their code (hidden)
    for i, char in enumerate(text):
        if char == "\x13":
            fields.append(True)
        elif char == "\x14":
            if fields:
                fields[-1] = False
        elif char == "\x15":
            if fields:
                fields.pop()
        elif any(fields):
            continue
        elif char in "\r\x07\x0c":
            p = properties(i)
            depth = p.values.get(SPRM_DEPTH, 1 if p.values.get(SPRM_IN_TABLE) else 0)
            row_end = char != "\x0c" and p.values.get(SPRM_ROW_END) and depth <= 1
            if char == "\x07" and depth <= 1 and not row_end:
                row.append(one_line("".join(current)))  # the end of a cell
                current = []
            elif row_end:
                table.append((row, p.edges))
                row, current = [], []
            elif depth >= 1:
                current.append(" ")  # a paragraph, a page break or a nested table's mark in a cell
            else:
                end_paragraph()  # a paragraph mark, or a page or section break
        elif char == "\x0b":  # a line break
            current.append(" ")
        elif char in "\x01\x08\x1f":  # pictures, drawn objects, optional hyphens
            continue
        elif char == "\x1e":  # non-breaking hyphen
            current.append("-")
        else:
            current.append(char)
    end_paragraph()
    return "\n\n".join(b for b in blocks if b) + "\n"


# The converters of each file type, with the extension of the text they give: Markdown, or
# a PDF's plain text.
CONVERTERS: dict[str, tuple[Callable[[Path], str], str]] = {
    ".xml": (jats, ".md"),
    ".nxml": (jats, ".md"),
    ".html": (html, ".md"),
    ".htm": (html, ".md"),
    ".pdf": (pdf, ".txt"),
    ".docx": (word, ".md"),
    ".doc": (doc, ".md"),
}
# Spreadsheets: each sheet's name and rows, one TSV text per sheet.
SPREADSHEETS: dict[str, Callable[[Path], list[tuple[str, str]]]] = {
    ".xlsx": xlsx,
    ".xlsm": xlsx,
    ".xls": xls,
}
# Already text: cleaned in place.
AS_THEY_ARE = {".csv", ".tsv", ".txt"}
# The encodings of text files other than UTF-8: spreadsheet exports (Excel's "Unicode text",
# Windows, old Mac).
ENCODINGS = ["utf_16", "utf_16_le", "utf_16_be", "cp1252", "mac_roman"]


def _read[T](converter: Callable[[Path], T], path: Path) -> T:
    try:
        return converter(path)
    except ParseError:
        raise
    except Exception as e:  # a damaged or encrypted file, from anywhere: refused with why
        raise ParseError(f"cannot be read: {type(e).__name__}: {e}") from e


def _visible(text: str) -> str:
    # Zero-width spaces and soft hyphens: invisible, and tokens all the same.
    return re.sub("[\u200b\u200c\u200d\u00ad\ufeff]", "", text)


def parse(path: Path) -> str:
    """A full text or a document, as Markdown or plain text by its type."""
    found = CONVERTERS.get(path.suffix.lower())
    if found is None:
        raise ParseError(f"no converter for {path.suffix or 'files without extension'}")
    return _visible(_read(found[0], path))


def sheets(path: Path) -> list[tuple[str, str]]:
    """A spreadsheet's sheets, in workbook order: each its name and its rows as TSV."""
    return [(name, _visible(rows)) for name, rows in _read(SPREADSHEETS[path.suffix.lower()], path)]


def _decoded(raw: bytes) -> str:
    """A text file's text: UTF-8 (never with NUL bytes, which UTF-16 has), else in the encoding
    found among `ENCODINGS`."""
    if b"\x00" not in raw:
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            pass
    found = charset_normalizer.from_bytes(raw, cp_isolation=ENCODINGS).best()
    if found is None:
        raise ParseError("unknown encoding")
    return str(found)


def _csv(text: str, separator: str) -> list[list[str]]:
    return list(csv.reader(io.StringIO(text, newline=""), delimiter=separator))


def _fit(rows: list[list[str]]) -> tuple[int, int]:
    """How many rows share the commonest number of fields, and that number; none for one field."""
    widths = Counter(len(cells) for cells in rows if any(cells))
    width, count = max(widths.items(), key=lambda w: (w[1], w[0]), default=(0, 0))
    return (count, width) if width > 1 else (0, 0)


def _table(text: str, suffix: str) -> str:
    """A CSV or TSV file's text, its rows padded as a sheet's: a TSV split at its tabs, a CSV
    read as CSV, its separator the one, among comma, semicolon and tab, that gives the most rows
    the same number of fields (left as it is when each gives one field). Plain text as it is."""
    if suffix == ".tsv":
        rows = _padded(line.split("\t") for line in text.splitlines())
        return "".join("\t".join(cells) + "\n" for cells in rows)
    if suffix != ".csv":
        return text
    try:
        separator, rows = max(((s, _csv(text, s)) for s in ",;\t"), key=lambda s: _fit(s[1]))
    except csv.Error:  # a field over the csv module's size limit
        return text
    if _fit(rows) == (0, 0):
        return text
    out = io.StringIO()
    csv.writer(out, delimiter=separator, lineterminator="\n").writerows(_padded(rows))
    return out.getvalue()


def _write(path: Path, text: str) -> int:
    data = text.encode()
    path.write_bytes(data)
    return len(data)


def convert(base: Path, file: str) -> dict[str, Any]:
    """A file, at a path relative to `base`, converted to its texts, written beside it; returns
    its entry of the manifest, paths relative to `base`: `file`, then `text` and `bytes`, or for
    a spreadsheet `sheets`, each with `sheet` (its name), `text` and `bytes`. CSV, TSV and plain
    text are their own text, rewritten in UTF-8 without invisible characters, a CSV's or TSV's
    rows padded, when they were not. A file that cannot be converted writes nothing."""
    path = base / file
    suffix = path.suffix.lower()
    if not path.is_file():
        raise ParseError("no such file")
    if suffix in AS_THEY_ARE:
        raw = path.read_bytes()
        cleaned = _table(_visible(_decoded(raw)), suffix).encode()
        if cleaned != raw:
            path.write_bytes(cleaned)
        return {"file": file, "text": file, "bytes": len(cleaned)}
    if suffix in SPREADSHEETS:
        found: list[dict[str, Any]] = []
        for n, (name, rows) in enumerate(sheets(path), 1):
            text = f"{file}.{n}.tsv"
            found.append({"sheet": name, "text": text, "bytes": _write(base / text, rows)})
        return {"file": file, "sheets": found}
    content = parse(path)
    text = file + CONVERTERS[suffix][1]
    return {"file": file, "text": text, "bytes": _write(base / text, content)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert full texts or supplements to text.")
    parser.add_argument("dir", type=Path)
    parser.add_argument("files", nargs="+", metavar="file")
    args = parser.parse_args()
    refused = False
    for file in args.files:
        try:
            print(json.dumps(convert(args.dir, file), ensure_ascii=False))
        except ParseError as e:
            print(f"{file}: not converted, {e}", file=sys.stderr)
            refused = True
    if refused:
        sys.exit(1)
